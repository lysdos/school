"""在现有 xlsx 基础上追加『班级』列,默认给空。

只追加列，不改原数据，方便迁移。运行:
    python tools/update_roster_with_class.py
"""
from __future__ import annotations

from pathlib import Path

import openpyxl

XLSX = Path("data/学生记录.xlsx")


def main() -> int:
    if not XLSX.exists():
        print(f"[!] 找不到 {XLSX}")
        return 1
    wb = openpyxl.load_workbook(XLSX)
    ws = wb.active
    header = [c.value for c in ws[1]]
    if len(header) >= 4 and header[3] == "班级":
        print("[i] 已是 4 列格式,跳过")
        return 0
    ws.cell(row=1, column=4, value="班级")
    for r in range(2, ws.max_row + 1):
        ws.cell(row=r, column=4, value=None)
    wb.save(XLSX)
    print(f"[ok] 升级 {XLSX} -> 4 列,班级为空,可后台或 Excel 手动填")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
