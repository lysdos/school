"""从 xlsx 导入学生名单。幂等：重复导入不会产生重复行，也不会清掉运行状态。"""

from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook

from .db import Database
from .models import normalize_epc

# 历史 3 列格式仍然兼容: 没有班级列时 class_name = None
# 新格式推荐 4 列: EPC / 学号 / 姓名 / 班级
EXPECTED_HEADERS_3 = ("EPC", "学号", "姓名")
EXPECTED_HEADERS_4 = ("EPC", "学号", "姓名", "班级")


def _cell_text(value: Any) -> str:
    """Excel 里学号常被存成浮点（1424764951.0），这里还原成整数文本。"""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def import_roster(db: Database, xlsx_path: Path) -> int:
    path = Path(xlsx_path)
    if not path.exists():
        raise FileNotFoundError(f"名单文件不存在: {path}")

    workbook = load_workbook(path, data_only=True, read_only=True)
    try:
        sheet = workbook.active
        rows: Iterable[tuple] = sheet.iter_rows(values_only=True)

        try:
            raw_header = tuple(_cell_text(c) for c in next(rows))
        except StopIteration:
            raise ValueError(f"名单文件为空: {path}") from None

        if raw_header[:4] == EXPECTED_HEADERS_4:
            header_len = 4
        elif raw_header[:3] == EXPECTED_HEADERS_3:
            header_len = 3
        else:
            raise ValueError(
                f"表头不符，期望 {EXPECTED_HEADERS_4} 或 {EXPECTED_HEADERS_3}，"
                f"实际 {raw_header}"
            )

        created_at = datetime.now().isoformat(timespec="seconds")
        count = 0
        for row in rows:
            cells = list(row[:header_len]) + [None] * (header_len - len(row[:header_len]))
            epc = normalize_epc(_cell_text(cells[0]))
            student_no = _cell_text(cells[1])
            name = _cell_text(cells[2])
            class_name = _cell_text(cells[3]) if header_len == 4 else None
            if not epc or not student_no or not name:
                continue
            db.upsert_student(epc, student_no, name, created_at,
                              class_name=class_name or None)
            count += 1
        return count
    finally:
        workbook.close()
